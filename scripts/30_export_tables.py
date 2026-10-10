#!/usr/bin/env python3
"""30 — export the dataset summary and the supplementary data tables (xlsx + docx)."""
import os
ROOT = os.environ.get("BCBM_ROOT") or os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import subprocess, os, re
import numpy as np, pandas as pd
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

B=ROOT
OUT=os.path.join(ROOT, "out"); FIG=f"{OUT}/figdata"; TAB=os.path.join(ROOT, "tables"); os.makedirs(TAB,exist_ok=True)

# ---------- 取 GSE324453 样本元数据 ----------
raw=subprocess.run(["curl","-sL","-A","Mozilla/5.0","--max-time","90",
  "https://www.ncbi.nlm.nih.gov/geo/query/acc.cgi?acc=GSE324453&targ=gsm&form=text&view=brief"],
  capture_output=True,text=True).stdout.splitlines()
cur=None; md=[]
for l in raw:
    if l.startswith("!Sample_title"): cur={"title":l.split("= ",1)[1].strip()}
    elif l.startswith("!Sample_characteristics_ch1") and cur is not None:
        k,v=l.split("= ",1)[1].strip().split(": ",1); cur[k]=v
    elif l.startswith("^SAMPLE") and (cur is not None): md.append(cur); cur=None
if cur is not None: md.append(cur)
sm=pd.DataFrame(md)
sm["sample"]=sm["title"].str.replace("snRNA-Seq ","",regex=False)
sm["region"]=sm["bcbm location"].str.strip()

# ---------- dataset summary ----------
T1=pd.DataFrame([
 ["GSE322943","Region discovery (bulk)","Bulk RNA-seq","99","cerebrum 65 / cerebellum 34",
  "BCBM (all breast cancer)","Region-balanced across 3 batches; GC-content GHGA raw"],
 ["GSE324453","Cell-type localization (snRNA)","Single-nucleus RNA-seq (10x)","28 (56,268 nuclei)",
  "cerebrum 24 / cerebellum 4","BCBM (TNBC-focused)","Same cohort/superseries as GSE322943 (GSE324454)"],
 ["GSE325196","Confounding counter-example","Single-nucleus RNA-seq + spatial","23","cerebellum 3 / cerebral lobes 20",
  "Multi-primary (BC, melanoma, NSCLC, SCLC, sarcoma)","Brain region fully confounded with primary tumour type"],
],columns=["Dataset","Role","Assay","n samples","Brain-region composition","Primary tumours","Notes"])

# ---------- sample metadata ----------
S1=sm[["sample","region"]].copy()
S1.columns=["Sample","Brain region (bcbm location)"]
for c in ["age","gender","nuclei loaded"]:
    if c in sm.columns: S1[c.capitalize()]=sm[c].values
meta=pd.read_csv(f"{OUT}/GSE324453/meta.csv",index_col=0)
S1["Nuclei (this study)"]=S1["Sample"].map(meta.groupby("sampleDemuxIDstrict").size()).fillna(0).astype(int)
S1=S1.sort_values(["Brain region (bcbm location)","Sample"])

# ---------- bulk region DE ----------
de=pd.read_csv(f"{OUT}/03b_region_de_symbols.csv",index_col=0)
ad=pd.read_csv(f"{OUT}/03c_region_adjusted.csv",index_col=0)
S2=de.copy(); S2.index.name="Gene"
S2["log2FC_adj(brain-content)"]=ad.reindex(S2.index)["log2FC_adj"]
S2=S2[["log2FC","p","q","log2FC_adj(brain-content)"]].reset_index()

# ---------- snRNA compartment DE ----------
rows=[]
for f in sorted(os.listdir(OUT)):
    if f.startswith("11_GSE324453_region_"):
        c=f.replace("11_GSE324453_region_","").replace(".csv","")
        d=pd.read_csv(f"{OUT}/{f}",index_col=0)
        d=d.rename(columns={"lfc":"log2FC"}); d.insert(0,"Compartment",c)
        d.index.name="Gene"; rows.append(d.reset_index())
S3=pd.concat(rows,ignore_index=True)

# ---------- signature gene set ----------
feat=set(l.strip() for l in open(f"{OUT}/GSE324453/features.tsv"))
SET=de.query("q<0.05").query("log2FC<0").copy()
SET=SET[[g in feat for g in SET.index]]
S4=SET.reset_index().rename(columns={"index":"Gene"})
S4=S4[["Gene","log2FC","p","q"]]
print(f"datasets {T1.shape} | S1 {S1.shape} | S2 {S2.shape} | S3 {S3.shape} | S4 {S4.shape}")

# ---------- 写 xlsx ----------
with pd.ExcelWriter(f"{TAB}/datasets_summary.xlsx",engine="openpyxl") as w:
    T1.to_excel(w,index=False,sheet_name="datasets")
with pd.ExcelWriter(f"{TAB}/supplementary_tables.xlsx",engine="openpyxl") as w:
    S1.to_excel(w,index=False,sheet_name="S1_sample_meta")
    S2.to_excel(w,index=False,sheet_name="S2_bulk_region_DE")
    S3.to_excel(w,index=False,sheet_name="S3_snRNA_compartment_DE")
    S4.to_excel(w,index=False,sheet_name="S4_signature_106genes")

# ---------- 写 docx（三线表）----------
def _three_line(tbl):
    for i,row in enumerate(tbl.rows):
        for cell in row.cells:
            tcPr=cell._tc.get_or_add_tcPr()
            for tag in ("tcBorders","tcMar"):
                el=tcPr.find(qn("w:"+tag))
                if el is not None: tcPr.remove(el)
            bd=OxmlElement("w:tcBorders")
            for edge,on in [("top",i==0),("bottom",i==0 or i==len(tbl.rows)-1),("left",False),("right",False)]:
                e=OxmlElement("w:"+edge)
                e.set(qn("w:val"),"single" if on else "none")
                e.set(qn("w:sz"),"8"); e.set(qn("w:color"),"000000")
                bd.append(e)
            tcPr.append(bd)
def _font(doc):
    for p in doc.paragraphs+[p for t in doc.tables for r in t.rows for c in r.cells for p in c.paragraphs]:
        for run in p.runs:
            run.font.name="Times New Roman"; run.font.size=Pt(9)
            run._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
def add_table(doc,df,title=None):
    if title:
        _p = doc.add_paragraph(title)
        if _p.runs:
            _p.runs[0].bold = True
    t=doc.add_table(rows=1,cols=df.shape[1])
    for j,c in enumerate(df.columns): t.rows[0].cells[j].text=str(c)
    for _,row in df.iterrows():
        cells=t.add_row().cells
        for j,v in enumerate(row): cells[j].text="" if pd.isna(v) else (f"{v:.4g}" if isinstance(v,float) else str(v))
    _three_line(t); return t
doc=Document(); doc.add_paragraph("Datasets and cohorts")
add_table(doc,T1)
_font(doc); doc.save(f"{TAB}/datasets_summary.docx")
doc=Document()
for tt,df in [("GSE324453 sample metadata (n=28)",S1),
              ("Bulk brain-region DE genes (GSE322943), first 100 rows",S2[:100]),
              ("snRNA compartment-wise region DE (GSE324453), first 100 rows",S3[:100]),
              ("Cerebellum-up signature gene set (n=%d)"%len(S4),S4)]:
    add_table(doc,df,title=tt)
_font(doc); doc.save(f"{TAB}/supplementary_tables.docx")
print("落盘:",TAB)
